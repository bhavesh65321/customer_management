"""
GST API routes
==============

GET  /api/gst/settings              → Get store GST configuration
POST /api/gst/settings              → Update store GST settings  (manager+)
GET  /api/gst/calculate             → Preview GST before saving a bill (staff+)
GET  /api/gst/gstr1?month=&year=    → GSTR-1 report  (manager+)
GET  /api/gst/gstr3b?month=&year=   → GSTR-3B summary (manager+)
GET  /api/gst/hsn-summary?month=&year=  → HSN-wise breakup (manager+)
GET  /api/gst/export?month=&year=   → Download ZIP for CA (manager+)
"""

import calendar
import io
import zipfile
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import extract

from config.database import get_db
from dependencies import require_manager, require_staff, get_token_payload
from models.store import Store
from models.transactional import Transaction
from models.customer import Customer
from schemas.gst_schema import (
    GSTSettingsUpdate,
    GSTSettingsOut,
    GSTBreakdown,
    GSTR1Report,
    GSTR3BSummary,
    B2BInvoiceRow,
    B2CLargeRow,
    B2CSmallSummary,
    HSNRow,
)
from utils.gst_utils import (
    calculate_gst,
    get_financial_year,
    get_hsn_description,
    B2C_LARGE_THRESHOLD,
)
from utils.activity import log_activity

router = APIRouter(tags=["GST"])


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _get_store(store_id: int, db: Session) -> Store:
    store = db.query(Store).filter(Store.id == store_id).first()
    if not store:
        raise HTTPException(status_code=404, detail="Store not found")
    return store


def _resolve_store_id(payload: dict, db: Session) -> int:
    """Return store_id from JWT, or first active store for admin with no store."""
    store_id = payload.get("store_id")
    if store_id:
        return store_id
    role = payload.get("role", "staff")
    if role == "admin":
        first = db.query(Store).filter(Store.is_active == True).first()
        if first:
            return first.id
    raise HTTPException(status_code=403, detail="No store linked to your account")


def _period_label(month: int, year: int) -> str:
    return f"{calendar.month_name[month]} {year}"


def _gst_transactions(store_id: int, month: int, year: int, db: Session):
    """All GST-computed transactions for a store/month."""
    return (
        db.query(Transaction, Customer)
        .join(Customer, Transaction.customer_id == Customer.id, isouter=True)
        .filter(
            Transaction.store_id == store_id,
            Transaction.gst_computed == True,
            extract("month", Transaction.date) == month,
            extract("year", Transaction.date) == year,
        )
        .all()
    )


# ---------------------------------------------------------------------------
# Store GST settings
# ---------------------------------------------------------------------------

@router.get("/settings", response_model=GSTSettingsOut)
def get_gst_settings(
    db: Session = Depends(get_db),
    payload: dict = Depends(require_staff),
):
    """View current GST config for the store."""
    store_id = _resolve_store_id(payload, db)
    return _get_store(store_id, db)


@router.post("/settings", response_model=GSTSettingsOut)
def update_gst_settings(
    body: GSTSettingsUpdate,
    db: Session = Depends(get_db),
    payload: dict = Depends(require_manager),
):
    """
    Manager / admin configures store's GSTIN, state code,
    invoice prefix and default HSN codes.
    """
    store_id = _resolve_store_id(payload, db)
    store = _get_store(store_id, db)

    updates = body.model_dump(exclude_unset=True)

    # Auto-derive invoice prefix from store name if not explicitly set
    if "invoice_prefix" not in updates or not updates.get("invoice_prefix"):
        import re
        updates["invoice_prefix"] = re.sub(r"[^A-Za-z0-9]", "", store.name).upper()[:4]

    for key, value in updates.items():
        setattr(store, key, value)

    db.commit()
    db.refresh(store)
    log_activity(db, payload, action="updated", entity_type="gst_settings",
                 entity_id=str(store_id),
                 message=f"Updated GST settings for store #{store_id} (GSTIN={store.gstin or 'not set'})",
                 store_id=store_id)
    return store


# ---------------------------------------------------------------------------
# Live GST calculator (billing screen preview — no DB write)
# ---------------------------------------------------------------------------

@router.get("/calculate", response_model=GSTBreakdown)
def preview_gst(
    item_value: float = Query(..., ge=0, description="Jewellery item value (₹)"),
    making_charges: float = Query(0.0, ge=0, description="Making charges (₹)"),
    hsn_code: str = Query("7113", description="HSN for main item"),
    making_hsn_code: str = Query("9988", description="HSN for making charges"),
    is_interstate: bool = Query(False, description="True if customer is in a different state"),
    payload: dict = Depends(require_staff),
):
    """
    Preview full GST breakdown before saving a bill.
    Staff calls this live as they type on the billing screen.
    Returns CGST/SGST (intra-state) or IGST (inter-state) split.
    """
    result = calculate_gst(
        item_value=item_value,
        making_charges=making_charges,
        hsn_code=hsn_code,
        making_hsn_code=making_hsn_code,
        is_interstate=is_interstate,
    )
    return GSTBreakdown(
        hsn_code=result["hsn_code"],
        making_hsn_code=result["making_hsn_code"],
        tax_rate=result["tax_rate"],
        making_tax_rate=result["making_tax_rate"],
        taxable_value=result["taxable_value"],
        making_charges=result["making_charges"],
        cgst_amount=result["cgst_amount"],
        sgst_amount=result["sgst_amount"],
        igst_amount=result["igst_amount"],
        making_cgst=result["making_cgst"],
        making_sgst=result["making_sgst"],
        making_igst=result["making_igst"],
        is_interstate=result["is_interstate"],
        customer_gstin=None,
        invoice_number=None,
        total_tax=result["total_tax"],
        grand_total=result["grand_total"],
    )


# ---------------------------------------------------------------------------
# GSTR-1
# ---------------------------------------------------------------------------

@router.get("/gstr1", response_model=GSTR1Report)
def get_gstr1(
    month: int = Query(..., ge=1, le=12),
    year: int = Query(..., ge=2020),
    db: Session = Depends(get_db),
    payload: dict = Depends(require_manager),
):
    """
    GSTR-1 — Monthly outward supply report.
    Splits automatically into B2B, B2C Large (>₹2.5L), B2C Small, and HSN Summary.
    CA uses this to upload to GST portal.
    """
    store_id = _resolve_store_id(payload, db)
    store = _get_store(store_id, db)
    rows = _gst_transactions(store_id, month, year, db)

    b2b: list[B2BInvoiceRow] = []
    b2c_large: list[B2CLargeRow] = []
    b2c_small = dict(taxable=0.0, cgst=0.0, sgst=0.0, igst=0.0, invoice_total=0.0)
    hsn_map: dict[str, dict] = {}
    total_taxable = total_cgst = total_sgst = total_igst = 0.0

    for txn, customer in rows:
        inv_no = txn.invoice_number or f"INV-{txn.id}"
        inv_date = txn.date.strftime("%d-%m-%Y") if txn.date else ""
        cname = customer.name if customer else "Walk-in"
        item_amt = float(txn.taxable_value or 0) + float(txn.making_charges or 0)
        tax_amt = float(txn.cgst_amount or 0) + float(txn.sgst_amount or 0) + float(txn.igst_amount or 0)
        inv_total = round(item_amt + tax_amt, 2)

        total_taxable += float(txn.taxable_value or 0)
        total_cgst += float(txn.cgst_amount or 0)
        total_sgst += float(txn.sgst_amount or 0)
        total_igst += float(txn.igst_amount or 0)

        # Accumulate into HSN map
        for hsn, val, cgst, sgst, igst in [
            (
                txn.hsn_code,
                txn.taxable_value,
                float(txn.cgst_amount or 0) - float(txn.making_cgst or 0),
                float(txn.sgst_amount or 0) - float(txn.making_sgst or 0),
                float(txn.igst_amount or 0) - float(txn.making_igst or 0),
            ),
            (
                txn.making_hsn_code,
                txn.making_charges,
                float(txn.making_cgst or 0),
                float(txn.making_sgst or 0),
                float(txn.making_igst or 0),
            ),
        ]:
            if hsn and float(val or 0) > 0:
                if hsn not in hsn_map:
                    hsn_map[hsn] = {"qty": 0, "value": 0.0, "cgst": 0.0, "sgst": 0.0, "igst": 0.0}
                hsn_map[hsn]["qty"] += 1
                hsn_map[hsn]["value"] += float(val or 0)
                hsn_map[hsn]["cgst"] += cgst
                hsn_map[hsn]["sgst"] += sgst
                hsn_map[hsn]["igst"] += igst

        # Route to B2B / B2C Large / B2C Small
        if txn.customer_gstin:
            b2b.append(B2BInvoiceRow(
                invoice_number=inv_no,
                invoice_date=inv_date,
                customer_name=cname,
                customer_gstin=txn.customer_gstin,
                hsn_code=txn.hsn_code,
                taxable_value=round(float(txn.taxable_value or 0), 2),
                cgst=round(float(txn.cgst_amount or 0), 2),
                sgst=round(float(txn.sgst_amount or 0), 2),
                igst=round(float(txn.igst_amount or 0), 2),
                total_tax=round(tax_amt, 2),
                invoice_total=inv_total,
            ))
        elif inv_total >= B2C_LARGE_THRESHOLD:
            b2c_large.append(B2CLargeRow(
                invoice_number=inv_no,
                invoice_date=inv_date,
                customer_name=cname,
                hsn_code=txn.hsn_code,
                taxable_value=round(float(txn.taxable_value or 0), 2),
                cgst=round(float(txn.cgst_amount or 0), 2),
                sgst=round(float(txn.sgst_amount or 0), 2),
                igst=round(float(txn.igst_amount or 0), 2),
                invoice_total=inv_total,
            ))
        else:
            b2c_small["taxable"] += float(txn.taxable_value or 0)
            b2c_small["cgst"] += float(txn.cgst_amount or 0)
            b2c_small["sgst"] += float(txn.sgst_amount or 0)
            b2c_small["igst"] += float(txn.igst_amount or 0)
            b2c_small["invoice_total"] += inv_total

    hsn_rows = [
        HSNRow(
            hsn_code=hsn,
            description=get_hsn_description(hsn),
            total_quantity=round(d["qty"], 2),
            total_value=round(d["value"], 2),
            total_cgst=round(d["cgst"], 2),
            total_sgst=round(d["sgst"], 2),
            total_igst=round(d["igst"], 2),
            total_tax=round(d["cgst"] + d["sgst"] + d["igst"], 2),
        )
        for hsn, d in sorted(hsn_map.items())
    ]

    return GSTR1Report(
        store_name=store.name,
        gstin=store.gstin,
        month=month,
        year=year,
        period_label=_period_label(month, year),
        b2b=b2b,
        b2c_large=b2c_large,
        b2c_small=B2CSmallSummary(
            month=_period_label(month, year),
            total_taxable_value=round(b2c_small["taxable"], 2),
            total_cgst=round(b2c_small["cgst"], 2),
            total_sgst=round(b2c_small["sgst"], 2),
            total_igst=round(b2c_small["igst"], 2),
            total_tax=round(b2c_small["cgst"] + b2c_small["sgst"] + b2c_small["igst"], 2),
            total_invoice_value=round(b2c_small["invoice_total"], 2),
        ),
        hsn_summary=hsn_rows,
        total_taxable_value=round(total_taxable, 2),
        total_cgst=round(total_cgst, 2),
        total_sgst=round(total_sgst, 2),
        total_igst=round(total_igst, 2),
        total_tax_collected=round(total_cgst + total_sgst + total_igst, 2),
    )


# ---------------------------------------------------------------------------
# GSTR-3B
# ---------------------------------------------------------------------------

@router.get("/gstr3b", response_model=GSTR3BSummary)
def get_gstr3b(
    month: int = Query(..., ge=1, le=12),
    year: int = Query(..., ge=2020),
    db: Session = Depends(get_db),
    payload: dict = Depends(require_manager),
):
    """
    GSTR-3B — Monthly consolidated summary.
    CA uses this to fill the tax return on GST portal.
    """
    store_id = _resolve_store_id(payload, db)
    store = _get_store(store_id, db)
    rows = _gst_transactions(store_id, month, year, db)

    t = dict(sales=0.0, taxable=0.0, cgst=0.0, sgst=0.0, igst=0.0,
             b2b=0.0, b2c=0.0, interstate=0.0, intrastate=0.0)

    for txn, _ in rows:
        item_amt = float(txn.taxable_value or 0) + float(txn.making_charges or 0)
        tax_amt = float(txn.cgst_amount or 0) + float(txn.sgst_amount or 0) + float(txn.igst_amount or 0)
        inv_total = item_amt + tax_amt
        t["sales"] += inv_total
        t["taxable"] += item_amt
        t["cgst"] += float(txn.cgst_amount or 0)
        t["sgst"] += float(txn.sgst_amount or 0)
        t["igst"] += float(txn.igst_amount or 0)
        if txn.customer_gstin:
            t["b2b"] += inv_total
        else:
            t["b2c"] += inv_total
        if txn.is_interstate:
            t["interstate"] += inv_total
        else:
            t["intrastate"] += inv_total

    return GSTR3BSummary(
        store_name=store.name,
        gstin=store.gstin,
        month=month,
        year=year,
        period_label=_period_label(month, year),
        total_sales=round(t["sales"], 2),
        total_taxable_value=round(t["taxable"], 2),
        total_cgst=round(t["cgst"], 2),
        total_sgst=round(t["sgst"], 2),
        total_igst=round(t["igst"], 2),
        total_tax_payable=round(t["cgst"] + t["sgst"] + t["igst"], 2),
        b2b_sales=round(t["b2b"], 2),
        b2c_sales=round(t["b2c"], 2),
        interstate_sales=round(t["interstate"], 2),
        intrastate_sales=round(t["intrastate"], 2),
    )


# ---------------------------------------------------------------------------
# HSN Summary (standalone — useful for quarterly returns too)
# ---------------------------------------------------------------------------

@router.get("/hsn-summary", response_model=list[HSNRow])
def get_hsn_summary(
    month: int = Query(..., ge=1, le=12),
    year: int = Query(..., ge=2020),
    db: Session = Depends(get_db),
    payload: dict = Depends(require_manager),
):
    """HSN-wise tax breakup for a given month — needed in both GSTR-1 and annual return."""
    store_id = _resolve_store_id(payload, db)
    rows = _gst_transactions(store_id, month, year, db)
    hsn_map: dict[str, dict] = {}

    for txn, _ in rows:
        for hsn, val, cgst, sgst, igst in [
            (txn.hsn_code, txn.taxable_value,
             float(txn.cgst_amount or 0) - float(txn.making_cgst or 0),
             float(txn.sgst_amount or 0) - float(txn.making_sgst or 0),
             float(txn.igst_amount or 0) - float(txn.making_igst or 0)),
            (txn.making_hsn_code, txn.making_charges,
             float(txn.making_cgst or 0), float(txn.making_sgst or 0), float(txn.making_igst or 0)),
        ]:
            if hsn and float(val or 0) > 0:
                if hsn not in hsn_map:
                    hsn_map[hsn] = {"qty": 0, "value": 0.0, "cgst": 0.0, "sgst": 0.0, "igst": 0.0}
                hsn_map[hsn]["qty"] += 1
                hsn_map[hsn]["value"] += float(val or 0)
                hsn_map[hsn]["cgst"] += cgst
                hsn_map[hsn]["sgst"] += sgst
                hsn_map[hsn]["igst"] += igst

    return [
        HSNRow(
            hsn_code=hsn,
            description=get_hsn_description(hsn),
            total_quantity=round(d["qty"], 2),
            total_value=round(d["value"], 2),
            total_cgst=round(d["cgst"], 2),
            total_sgst=round(d["sgst"], 2),
            total_igst=round(d["igst"], 2),
            total_tax=round(d["cgst"] + d["sgst"] + d["igst"], 2),
        )
        for hsn, d in sorted(hsn_map.items())
    ]


# ---------------------------------------------------------------------------
# CA Export — Excel ZIP  (one click → CA gets everything)
# ---------------------------------------------------------------------------

@router.get("/export")
def export_for_ca(
    month: int = Query(..., ge=1, le=12),
    year: int = Query(..., ge=2020),
    db: Session = Depends(get_db),
    payload: dict = Depends(require_manager),
):
    """
    One-click CA export package.
    Downloads a ZIP file containing:
      • GSTR1_{store}_{month}_{year}.xlsx  — upload to GST portal
      • GSTR3B_{store}_{month}_{year}.xlsx — CA fills tax return
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl not installed. Run: pip install openpyxl",
        )

    store_id = _resolve_store_id(payload, db)
    store = _get_store(store_id, db)
    rows = _gst_transactions(store_id, month, year, db)
    period = _period_label(month, year)

    HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
    HEADER_FONT = Font(bold=True, color="FFFFFF")

    def _header(ws, headers):
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=col, value=h)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
            c.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 20

    # ── GSTR-1 workbook ──────────────────────────────────────────────────
    wb1 = openpyxl.Workbook()

    ws_b2b = wb1.active
    ws_b2b.title = "B2B Invoices"
    _header(ws_b2b, ["Invoice No", "Date", "Customer", "GSTIN", "HSN",
                      "Taxable Value", "CGST", "SGST", "IGST", "Total Tax", "Invoice Total"])

    ws_large = wb1.create_sheet("B2C Large (>2.5L)")
    _header(ws_large, ["Invoice No", "Date", "Customer", "HSN",
                        "Taxable Value", "CGST", "SGST", "IGST", "Invoice Total"])

    ws_small = wb1.create_sheet("B2C Small (Consolidated)")
    _header(ws_small, ["Period", "Taxable Value", "CGST", "SGST", "IGST", "Total Tax", "Invoice Total"])

    ws_hsn = wb1.create_sheet("HSN Summary")
    _header(ws_hsn, ["HSN Code", "Description", "Total Items",
                      "Taxable Value", "CGST", "SGST", "IGST", "Total Tax"])

    b2c_small = dict(taxable=0.0, cgst=0.0, sgst=0.0, igst=0.0, inv_total=0.0)
    hsn_map: dict[str, dict] = {}

    for txn, customer in rows:
        cname = customer.name if customer else "Walk-in"
        inv_no = txn.invoice_number or f"INV-{txn.id}"
        inv_date = txn.date.strftime("%d-%m-%Y") if txn.date else ""
        item_amt = float(txn.taxable_value or 0) + float(txn.making_charges or 0)
        tax_amt = (float(txn.cgst_amount or 0) + float(txn.sgst_amount or 0)
                   + float(txn.igst_amount or 0))
        inv_total = round(item_amt + tax_amt, 2)

        if txn.customer_gstin:
            ws_b2b.append([
                inv_no, inv_date, cname, txn.customer_gstin, txn.hsn_code,
                round(float(txn.taxable_value or 0), 2),
                round(float(txn.cgst_amount or 0), 2),
                round(float(txn.sgst_amount or 0), 2),
                round(float(txn.igst_amount or 0), 2),
                round(tax_amt, 2), inv_total,
            ])
        elif inv_total >= B2C_LARGE_THRESHOLD:
            ws_large.append([
                inv_no, inv_date, cname, txn.hsn_code,
                round(float(txn.taxable_value or 0), 2),
                round(float(txn.cgst_amount or 0), 2),
                round(float(txn.sgst_amount or 0), 2),
                round(float(txn.igst_amount or 0), 2),
                inv_total,
            ])
        else:
            b2c_small["taxable"] += float(txn.taxable_value or 0)
            b2c_small["cgst"] += float(txn.cgst_amount or 0)
            b2c_small["sgst"] += float(txn.sgst_amount or 0)
            b2c_small["igst"] += float(txn.igst_amount or 0)
            b2c_small["inv_total"] += inv_total

        for hsn, val in [(txn.hsn_code, txn.taxable_value),
                          (txn.making_hsn_code, txn.making_charges)]:
            if hsn and float(val or 0) > 0:
                if hsn not in hsn_map:
                    hsn_map[hsn] = {"qty": 0, "value": 0.0}
                hsn_map[hsn]["qty"] += 1
                hsn_map[hsn]["value"] += float(val or 0)

    ws_small.append([
        period,
        round(b2c_small["taxable"], 2),
        round(b2c_small["cgst"], 2),
        round(b2c_small["sgst"], 2),
        round(b2c_small["igst"], 2),
        round(b2c_small["cgst"] + b2c_small["sgst"] + b2c_small["igst"], 2),
        round(b2c_small["inv_total"], 2),
    ])

    for hsn, d in sorted(hsn_map.items()):
        ws_hsn.append([hsn, get_hsn_description(hsn), d["qty"], round(d["value"], 2),
                        "-", "-", "-", "-"])  # CGSt per HSN needs separate accumulation

    # ── GSTR-3B workbook ─────────────────────────────────────────────────
    wb2 = openpyxl.Workbook()
    ws3b = wb2.active
    ws3b.title = "GSTR-3B Summary"
    _header(ws3b, ["Description", "Amount (₹)"])

    t = dict(sales=0.0, taxable=0.0, cgst=0.0, sgst=0.0, igst=0.0)
    for txn, _ in rows:
        it = float(txn.taxable_value or 0) + float(txn.making_charges or 0)
        tx = (float(txn.cgst_amount or 0) + float(txn.sgst_amount or 0)
              + float(txn.igst_amount or 0))
        t["sales"] += it + tx
        t["taxable"] += it
        t["cgst"] += float(txn.cgst_amount or 0)
        t["sgst"] += float(txn.sgst_amount or 0)
        t["igst"] += float(txn.igst_amount or 0)

    ws3b.append(["Store Name", store.name])
    ws3b.append(["GSTIN", store.gstin or "Not configured"])
    ws3b.append(["State Code", store.state_code or "Not configured"])
    ws3b.append(["Period", period])
    ws3b.append([])
    ws3b.append(["Total Sales (incl. tax)", round(t["sales"], 2)])
    ws3b.append(["Total Taxable Value", round(t["taxable"], 2)])
    ws3b.append([])
    ws3b.append(["CGST Collected", round(t["cgst"], 2)])
    ws3b.append(["SGST Collected", round(t["sgst"], 2)])
    ws3b.append(["IGST Collected", round(t["igst"], 2)])
    ws3b.append(["TOTAL TAX PAYABLE", round(t["cgst"] + t["sgst"] + t["igst"], 2)])

    # Style the "TOTAL TAX PAYABLE" row
    last_row = ws3b.max_row
    for col in range(1, 3):
        ws3b.cell(last_row, col).font = Font(bold=True, color="CC0000")

    # ── Pack into ZIP ─────────────────────────────────────────────────────
    prefix = f"{store.name.replace(' ', '_')}_{calendar.month_abbr[month]}_{year}"
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for wb, fname in [
            (wb1, f"GSTR1_{prefix}.xlsx"),
            (wb2, f"GSTR3B_{prefix}.xlsx"),
        ]:
            buf = io.BytesIO()
            wb.save(buf)
            zf.writestr(fname, buf.getvalue())

    zip_buf.seek(0)
    return StreamingResponse(
        zip_buf,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{prefix}_GST_Package.zip"'},
    )
