import logging
from datetime import datetime, timedelta
import secrets
from io import BytesIO

from fastapi import APIRouter, HTTPException, Depends, Query, UploadFile
from fastapi.responses import StreamingResponse, JSONResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_
from models.customer import Customer
from models.customer_invite import CustomerInvite
from schemas.customer import CustomerCreate, CustomerResponse
from config.database import get_db
from dependencies import require_staff, require_manager, get_token_payload
from typing import List, Optional
from utils.activity import log_activity
from core.db_filters import resolve_write_store_id


router = APIRouter(tags=["Customers"])
logger = logging.getLogger(__name__)


def _normalize_phone(s: str) -> str:
    if not s:
        return ""
    return "".join(c for c in str(s).strip() if c.isdigit())


def _store_filter(q, payload):
    """Filter customers to the requesting user's store. Admins see all."""
    from sqlalchemy import false as sql_false
    store_id = payload.get("store_id")
    role = payload.get("role", "staff")
    if store_id is not None:
        return q.filter(Customer.store_id == store_id)
    if role == "admin":
        return q  # admin sees all customers
    return q.filter(sql_false())  # misconfigured account → see nothing


@router.post("/add", response_model=CustomerResponse, status_code=201, include_in_schema=True)
def add_customer(
    customer: CustomerCreate,
    db: Session = Depends(get_db),
    payload: dict = Depends(require_staff),
):
    try:
        store_id = resolve_write_store_id(payload, db)
        norm_phone = _normalize_phone(customer.primary_phone)
        if norm_phone:
            q = db.query(Customer).filter(Customer.store_id == store_id, Customer.is_active == True)
            for c in q.all():
                if _normalize_phone(c.primary_phone) == norm_phone:
                    raise HTTPException(status_code=400, detail="A customer with this phone number already exists in your store")
        if customer.email and customer.email.strip():
            existing_email = db.query(Customer).filter(
                Customer.store_id == store_id,
                Customer.email == customer.email.strip(),
                Customer.is_active == True,
            ).first()
            if existing_email:
                raise HTTPException(status_code=400, detail="A customer with this email already exists in your store")
        new_customer = Customer(
            name=customer.name,
            father_name=customer.father_name,
            primary_phone=customer.primary_phone,
            secondary_phone=customer.secondary_phone,
            address=customer.address,
            city=customer.city,
            pincode=customer.pincode,
            gender=customer.gender,
            country=customer.country,
            email=customer.email,
            store_id=store_id,
        )
        db.add(new_customer)
        db.commit()
        db.refresh(new_customer)
        log_activity(
            db,
            payload,
            action="created",
            entity_type="customer",
            entity_id=str(new_customer.id),
            message=f"Added customer: {new_customer.name}",
        )
        return new_customer
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        logger.exception("add_customer failed")
        raise HTTPException(status_code=500, detail=str(e))


def _cell_value(cell):
    if cell is None:
        return None
    v = cell.value
    if v is None:
        return None
    return str(v).strip() if isinstance(v, str) else v


@router.post("/import")
def import_customers_excel(
    file: UploadFile,
    db: Session = Depends(get_db),
    payload: dict = Depends(require_manager),  # ⚠️ manager+ only
):
    store_id = resolve_write_store_id(payload, db)
    if not file.filename or not (file.filename.endswith(".xlsx") or file.filename.endswith(".xls")):
        raise HTTPException(status_code=400, detail="Upload an Excel file (.xlsx)")
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(status_code=500, detail="Excel support not installed (openpyxl)")
    try:
        wb = load_workbook(filename=file.file, read_only=True, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 2000)))
        wb.close()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid Excel file: {str(e)}")
    if not rows:
        return {"imported": 0, "errors": ["Sheet is empty."]}
    header = [_cell_value(c) for c in rows[0]]
    col_index = {}
    for i, h in enumerate(header):
        if h:
            key = str(h).lower().replace(" ", "_")
            if key not in col_index:
                col_index[key] = i
    name_col = col_index.get("name") if "name" in col_index else None
    phone_col = col_index.get("phone") or col_index.get("primary_phone") or col_index.get("mobile")
    if name_col is None or phone_col is None:
        return {
            "imported": 0,
            "errors": ["Excel must have columns 'Name' and 'Phone' (or 'Primary Phone' / 'Mobile')."],
        }
    email_col = col_index.get("email")
    address_col = col_index.get("address")
    father_col = col_index.get("father_name") or col_index.get("father") or col_index.get("father's_name")
    city_col = col_index.get("city")
    pin_col = col_index.get("pincode") or col_index.get("postal_code") or col_index.get("zip")
    gender_col = col_index.get("gender")
    country_col = col_index.get("country")
    secondary_col = col_index.get("secondary_phone") or col_index.get("alt_phone") or col_index.get("alternate_phone")

    existing = (
        db.query(Customer.primary_phone, Customer.email)
        .filter(Customer.store_id == store_id, Customer.is_active == True)
        .all()
    )
    seen_phones = set()
    seen_emails = set()
    for p, em in existing:
        np = _normalize_phone(p or "")
        if np:
            seen_phones.add(np)
        if em and str(em).strip():
            seen_emails.add(str(em).strip().lower())

    imported = 0
    errors = []
    for row_idx, row in enumerate(rows[1:], start=2):
        cells = [_cell_value(c) for c in row]
        name = cells[name_col] if name_col < len(cells) else None
        phone = cells[phone_col] if phone_col < len(cells) else None
        if not name or not str(name).strip():
            errors.append(f"Row {row_idx}: Name is required")
            continue
        if not phone or not str(phone).strip():
            errors.append(f"Row {row_idx}: Phone is required")
            continue
        name = str(name).strip()
        phone = str(phone).strip()
        email = cells[email_col] if email_col is not None and email_col < len(cells) and cells[email_col] else None
        if email:
            email = str(email).strip()
        address = cells[address_col] if address_col is not None and address_col < len(cells) and cells[address_col] else None
        if address:
            address = str(address).strip()
        father_name = cells[father_col] if father_col is not None and father_col < len(cells) and cells[father_col] else None
        if father_name:
            father_name = str(father_name).strip() or None
        city = cells[city_col] if city_col is not None and city_col < len(cells) and cells[city_col] else None
        if city:
            city = str(city).strip() or None
        pincode = cells[pin_col] if pin_col is not None and pin_col < len(cells) and cells[pin_col] else None
        if pincode is not None:
            pincode = str(pincode).strip() or None
        gender = cells[gender_col] if gender_col is not None and gender_col < len(cells) and cells[gender_col] else None
        if gender:
            gender = str(gender).strip() or None
        country = cells[country_col] if country_col is not None and country_col < len(cells) and cells[country_col] else None
        if country:
            country = str(country).strip() or None
        secondary_phone = cells[secondary_col] if secondary_col is not None and secondary_col < len(cells) and cells[secondary_col] else None
        if secondary_phone:
            secondary_phone = str(secondary_phone).strip() or None

        norm_phone = _normalize_phone(phone)
        if not norm_phone:
            errors.append(f"Row {row_idx}: Invalid phone")
            continue
        if norm_phone in seen_phones:
            errors.append(f"Row {row_idx}: Phone already exists in store")
            continue
        email_lower = email.strip().lower() if email and email.strip() else None
        if email_lower and email_lower in seen_emails:
            errors.append(f"Row {row_idx}: Email already exists in store")
            continue
        try:
            new_customer = Customer(
                name=name,
                father_name=father_name,
                primary_phone=phone,
                secondary_phone=secondary_phone,
                email=email or None,
                address=address or None,
                city=city,
                pincode=pincode,
                gender=gender,
                country=country,
                store_id=store_id,
            )
            db.add(new_customer)
            db.commit()
            db.refresh(new_customer)
            imported += 1
            seen_phones.add(norm_phone)
            if email_lower:
                seen_emails.add(email_lower)
        except Exception as e:
            db.rollback()
            errors.append(f"Row {row_idx}: {str(e)}")
    return {"imported": imported, "errors": errors[:50]}


@router.put("/update/{customer_id}", response_model=CustomerResponse, include_in_schema=True)
def update_customer(
    customer_id: int,
    updated_data: CustomerCreate,
    db: Session = Depends(get_db),
    payload: dict = Depends(require_staff),
):
    try:
        q = db.query(Customer).filter(Customer.id == customer_id)
        q = _store_filter(q, payload)
        customer = q.first()
        if not customer:
            raise HTTPException(status_code=404, detail="Customer not found")
        norm_phone = _normalize_phone(updated_data.primary_phone)
        if norm_phone:
            q = db.query(Customer).filter(
                Customer.store_id == payload.get("store_id"),
                Customer.is_active == True,
                Customer.id != customer_id,
            )
            for c in q.all():
                if _normalize_phone(c.primary_phone) == norm_phone:
                    raise HTTPException(status_code=400, detail="Another customer in your store already has this phone number")
        if updated_data.email and updated_data.email.strip():
            existing_email = db.query(Customer).filter(
                Customer.store_id == payload.get("store_id"),
                Customer.email == updated_data.email.strip(),
                Customer.is_active == True,
                Customer.id != customer_id,
            ).first()
            if existing_email:
                raise HTTPException(status_code=400, detail="Another customer in your store already has this email")
        customer.name = updated_data.name
        customer.father_name = updated_data.father_name
        customer.primary_phone = updated_data.primary_phone
        customer.secondary_phone = updated_data.secondary_phone
        customer.address = updated_data.address
        customer.city = updated_data.city
        customer.pincode = updated_data.pincode
        customer.gender = updated_data.gender
        customer.country = updated_data.country
        customer.email = updated_data.email

        db.commit()
        db.refresh(customer)
        log_activity(
            db,
            payload,
            action="updated",
            entity_type="customer",
            entity_id=str(customer_id),
            message=f"Updated customer: {customer.name}",
        )
        return customer

    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        logger.exception("update_customer failed")
        raise HTTPException(status_code=500, detail=str(e))
    


@router.get("/search", response_model=List[CustomerResponse], include_in_schema=True)
def search_customers(
    query: str = "",
    db: Session = Depends(get_db),
    payload: dict = Depends(require_staff),
):
    q = db.query(Customer).filter(Customer.name.ilike(f"%{query}%"))
    q = _store_filter(q, payload)
    return q.all()


@router.get("/all", response_model=List[CustomerResponse], include_in_schema=True)
def get_all_customers(
    status: Optional[str] = Query(None, description="active, inactive, or all"),
    db: Session = Depends(get_db),
    payload: dict = Depends(require_staff),
):
    q = db.query(Customer)
    q = _store_filter(q, payload)
    if status == "inactive":
        q = q.filter(Customer.is_active == False)
    elif status != "all":
        q = q.filter(Customer.is_active == True)
    return q.all()


@router.get("/list", response_model=List[CustomerResponse])
def list_customers(
    search: Optional[str] = Query(None),
    status: Optional[str] = Query(None, description="active, inactive, or all"),
    sort: Optional[str] = Query(None),
    # ── Pagination (BE-01) ─────────────────────────────────────────────────
    # page=0 (default) keeps legacy behaviour — returns ALL records so
    # existing frontend calls without pagination still work unchanged.
    page: int = Query(0, ge=0, description="1-based page number. 0 = return all (legacy)"),
    page_size: int = Query(50, ge=1, le=500, description="Items per page (max 500)"),
    db: Session = Depends(get_db),
    payload: dict = Depends(require_staff),
):
    query = db.query(Customer)
    query = _store_filter(query, payload)

    if status == "inactive":
        query = query.filter(Customer.is_active == False)
    elif status != "all":
        query = query.filter(Customer.is_active == True)

    term = (search or "").strip()
    if term:
        like = f"%{term}%"
        query = query.filter(
            or_(
                Customer.name.ilike(like),
                Customer.primary_phone.ilike(like),
                Customer.city.ilike(like),
                Customer.email.ilike(like),
            )
        )

    if sort == "name":
        query = query.order_by(Customer.name.asc())
    elif sort == "recent":
        query = query.order_by(Customer.id.desc())
    elif sort == "oldest":
        query = query.order_by(Customer.id.asc())
    else:
        query = query.order_by(Customer.id.desc())

    total = query.count()

    if page > 0:
        # Paginated mode
        items = query.offset((page - 1) * page_size).limit(page_size).all()
        import math
        headers = {
            "X-Total-Count": str(total),
            "X-Page": str(page),
            "X-Page-Size": str(page_size),
            "X-Total-Pages": str(math.ceil(total / page_size) if total else 0),
            "Access-Control-Expose-Headers": "X-Total-Count, X-Page, X-Page-Size, X-Total-Pages",
        }
        return JSONResponse(
            content=[item.model_dump() if hasattr(item, "model_dump") else
                     {c.name: getattr(item, c.name) for c in item.__table__.columns}
                     for item in items],
            headers=headers,
        )

    # Legacy mode (page=0): return all, same as before
    return query.all()


@router.get("/import-template")
def download_customer_import_template():
    try:
        from openpyxl import Workbook
    except ImportError:
        raise HTTPException(status_code=500, detail="Excel support not installed (openpyxl)")
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    headers = [
        "Name",
        "Phone",
        "Email",
        "Address",
        "Father Name",
        "City",
        "Pincode",
        "Gender",
        "Country",
        "Secondary Phone",
    ]
    ws.append(headers)
    ws.append(
        [
            "Example Customer",
            "9876543210",
            "customer@example.com",
            "123 Main Street",
            "Father name",
            "Mumbai",
            "400001",
            "Male",
            "India",
            "9876543211",
        ]
    )
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="customers-import-template.xlsx"',
        },
    )


@router.post("/invite/{customer_id}")
def create_invite(
    customer_id: int,
    db: Session = Depends(get_db),
    payload: dict = Depends(require_manager),  # ⚠️ manager+ only
):
    q = db.query(Customer).filter(Customer.id == customer_id)
    q = _store_filter(q, payload)
    customer = q.first()
    if not customer:
        raise HTTPException(status_code=404, detail="Customer not found")
    db.query(CustomerInvite).filter(
        CustomerInvite.customer_id == customer_id,
        CustomerInvite.used_at.is_(None),
        CustomerInvite.expires_at > datetime.utcnow(),
    ).delete(synchronize_session=False)
    token = secrets.token_urlsafe(32)
    expires_at = datetime.utcnow() + timedelta(days=7)
    invite = CustomerInvite(
        customer_id=customer_id,
        token=token,
        expires_at=expires_at,
    )
    db.add(invite)
    db.commit()
    link = f"/customer/join?token={token}"
    return {"invite_link": link, "token": token, "expires_at": expires_at.isoformat()}


@router.get("/{customer_id}", response_model=CustomerResponse)
def get_customer(
    customer_id: int,
    db: Session = Depends(get_db),
    payload: dict = Depends(require_staff),
):
    q = db.query(Customer).filter(Customer.id == customer_id)
    q = _store_filter(q, payload)
    customer = q.first()
    if not customer:
        raise HTTPException(status_code=404, detail="Customer not found")
    return customer
    

@router.delete("/delete/{customer_id}")
def delete_customer(
    customer_id: int,
    db: Session = Depends(get_db),
    payload: dict = Depends(require_manager),  # ⚠️ manager+ only
):
    q = db.query(Customer).filter(Customer.id == customer_id)
    q = _store_filter(q, payload)
    customer = q.first()
    if not customer:
        raise HTTPException(status_code=404, detail="Customer not found")
    customer.is_active = False
    db.commit()
    return {"message": "Customer marked as inactive"}

          



